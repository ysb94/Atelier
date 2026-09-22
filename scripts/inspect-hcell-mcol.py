import openpyxl

wb = openpyxl.load_workbook(
    r"node_modules/.tmp/sabangnet-official.xlsx",
    read_only=True,
    data_only=True,
)
ws = wb[wb.sheetnames[0]]
shown = 0
for row in ws.iter_rows(min_row=4, max_col=3, values_only=True):
    code, name, mlist = row[0], row[1], row[2]
    if str(code) in {"100050", "100008", "100416"}:
        print(repr(code), repr(name), repr(mlist), type(mlist))
        shown += 1
    if shown >= 3:
        break
